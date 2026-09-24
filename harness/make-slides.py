#!/usr/bin/env python3
"""Build the defense slide deck (PPTX) from the outline in docs/bao-cao/slide-bao-ve.md.

Output: docs/bao-cao/slide-bao-ve.pptx (16:9, 16 main + 6 backup slides).
Images are downscaled/compressed so the deck stays small enough to upload.
Run from the repo root: python3 harness/make-slides.py
Requires: python-pptx, Pillow.
"""

import io
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "bao-cao", "slide-bao-ve.pptx")
FIG = os.path.join(ROOT, "results", "figures")
SHOT = os.path.join(ROOT, "results", "screenshots-app-v1")
MASK = os.path.join(ROOT, "masking", "screens")

BLUE = RGBColor(0x44, 0x72, 0xC4)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x53, 0x8A, 0x35)
CODE_BG = RGBColor(0xF2, 0xF2, 0xF2)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def shrink(path, max_w=1400, quality=82):
    """Return a JPEG-compressed, downscaled stream of the image + its size."""
    img = Image.open(path).convert("RGB")
    if img.width > max_w:
        img = img.resize((max_w, int(img.height * max_w / img.width)), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=quality)
    buf.seek(0)
    return buf, img.width, img.height


def add_slide():
    return prs.slides.add_slide(BLANK)


def box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.text_frame.word_wrap = True
    return tb


def para(tf, text, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         first=False, space=6, font="Calibri"):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return p


def title_bar(slide, text, tag=None):
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(0.16))  # MSO_SHAPE.RECTANGLE
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    tb = box(slide, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.9))
    para(tb.text_frame, text, size=27, bold=True, first=True)
    if tag:
        tg = box(slide, Inches(10.6), Inches(0.06), Inches(2.6), Inches(0.4))
        para(tg.text_frame, tag, size=12, color=GREY, align=PP_ALIGN.RIGHT, first=True)


def picture(slide, path, left, top, width=None, height=None, compress=True):
    if compress:
        stream, w, h = shrink(path)
        if width and not height:
            height = Emu(int(width * h / w))
        return slide.shapes.add_picture(stream, left, top, width=width, height=height)
    return slide.shapes.add_picture(path, left, top, width=width, height=height)


def code_box(slide, left, top, width, height, lines, size=13):
    shp = slide.shapes.add_shape(1, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = CODE_BG
    shp.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.1)
    first = True
    for ln, color in lines:
        para(tf, ln, size=size, color=color, first=first, space=2, font="Consolas")
        first = False
    return shp


def big_number(slide, text, sub, top=Inches(5.55), color=BLUE):
    tb = box(slide, Inches(0.55), top, Inches(12.2), Inches(1.0))
    para(tb.text_frame, text, size=30, bold=True, color=color,
         align=PP_ALIGN.CENTER, first=True)
    if sub:
        tb2 = box(slide, Inches(0.55), top + Inches(0.75), Inches(12.2), Inches(0.6))
        para(tb2.text_frame, sub, size=15, color=GREY, align=PP_ALIGN.CENTER, first=True)


def bullets(slide, items, left=Inches(0.55), top=Inches(1.45), width=Inches(12.2),
            height=Inches(5.4), size=17):
    tb = box(slide, left, top, width, height)
    first = True
    for it in items:
        if isinstance(it, tuple):
            text, kw = it
        else:
            text, kw = it, {}
        para(tb.text_frame, text, size=kw.get("size", size),
             bold=kw.get("bold", False), color=kw.get("color", DARK),
             first=first, space=kw.get("space", 10))
        first = False
    return tb


# ---------------------------------------------------------------- Slide 1: title
s = add_slide()
bar = s.shapes.add_shape(1, 0, Inches(2.0), SW, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
tb = box(s, Inches(0.9), Inches(2.35), Inches(11.5), Inches(1.9))
para(tb.text_frame, "Ứng dụng Vision-Language Model trong Kiểm thử Tự động", size=32, bold=True, align=PP_ALIGN.CENTER, first=True)
para(tb.text_frame, "Giao diện Người dùng Phần mềm", size=32, bold=True, align=PP_ALIGN.CENTER)
para(tb.text_frame, "Khóa luận tốt nghiệp — UIT, ĐHQG TP.HCM", size=18, color=GREY, align=PP_ALIGN.CENTER, space=20)
tb = box(s, Inches(0.9), Inches(5.1), Inches(11.5), Inches(1.4))
para(tb.text_frame, "Sinh viên: [Họ tên — MSSV]", size=18, align=PP_ALIGN.CENTER, first=True)
para(tb.text_frame, "GVHD: [Họ tên GVHD]", size=18, align=PP_ALIGN.CENTER)
para(tb.text_frame, "[Ngày bảo vệ, 2026]", size=15, color=GREY, align=PP_ALIGN.CENTER)

# ------------------------------------------------- Slide 2: brittle locator problem
s = add_slide()
title_bar(s, "Vấn đề: test GUI dựa trên locator rất giòn")
code_box(s, Inches(0.55), Inches(1.5), Inches(6.1), Inches(1.5), [
    ("// Nút xóa = icon thứ 3 trong ô thứ 8 của hàng", GREY),
    (".product-table tbody tr:nth-child(1)", DARK),
    ("  td:nth-child(8) .icon-btn:nth-child(3)", RED),
])
tb = box(s, Inches(0.55), Inches(3.15), Inches(6.1), Inches(3.4))
para(tb.text_frame, "Locator gắn chặt vào cấu trúc DOM và nhãn văn bản.", size=17, first=True)
para(tb.text_frame, "Giao diện đổi cách trình bày (đảo cột, đổi nhãn nút) — chức năng không đổi — nhưng test gãy hàng loạt.", size=17)
para(tb.text_frame, "→ Chi phí bảo trì test là gánh nặng thường trực; nặng nhất ở thành phần không có nhãn ngữ nghĩa (icon SVG, canvas, div thuần).", size=17, color=RED)
picture(s, os.path.join(SHOT, "products-v0.png"), Inches(6.9), Inches(1.45), width=Inches(3.3))
tb = box(s, Inches(10.35), Inches(2.4), Inches(2.6), Inches(1.5))
para(tb.text_frame, "V0 (gốc): cột Actions ở vị trí thứ 8…", size=13, color=GREY, first=True)
picture(s, os.path.join(SHOT, "products-v2.png"), Inches(6.9), Inches(4.45), width=Inches(3.3))
tb = box(s, Inches(10.35), Inches(5.4), Inches(2.6), Inches(1.6))
para(tb.text_frame, "…V2: Actions nhảy lên cột 1 → selector trỏ sai ô", size=13, color=RED, first=True)

# ------------------------------------------------------- Slide 3: VLM alternative
s = add_slide()
title_bar(s, "Hướng mới: VLM “nhìn” giao diện, test viết bằng ngôn ngữ tự nhiên")
tb = box(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(0.5))
para(tb.text_frame, "Cùng một kịch bản A3 (thêm sản phẩm) trong hai bộ test đối xứng 1–1:", size=16, first=True)
code_box(s, Inches(0.55), Inches(1.95), Inches(6.0), Inches(2.6), [
    ("// Locator-based (Playwright)", GREY),
    ("page.click(\"xpath=//button[", DARK),
    ("  text()='Add product']\")", RED),
    ("modal.locator('label:nth-of-type(1)", DARK),
    ("  input').fill('Energy Drink 250ml')", DARK),
    ("expect(locator('.product-table tbody", DARK),
    ("  tr')).toHaveCount(13)", DARK),
], size=12)
code_box(s, Inches(6.8), Inches(1.95), Inches(6.0), Inches(2.6), [
    ("// VLM-based (Midscene.js + Qwen3-VL)", GREY),
    ("aiTap('the primary button in the", GREEN),
    ("  toolbar that adds a new product')", GREEN),
    ("aiInput('the Name field in the product", GREEN),
    ("  dialog', {value:'Energy Drink 250ml'})", GREEN),
    ("aiNumber('how many product rows does", GREEN),
    ("  the table contain?')  // == 13", GREEN),
], size=12)
tb = box(s, Inches(0.55), Inches(4.85), Inches(12.2), Inches(1.9))
para(tb.text_frame, "VLM đọc screenshot và thao tác theo ngữ nghĩa thị giác → về nguyên tắc miễn nhiễm với thay đổi trình bày.", size=17, first=True)
para(tb.text_frame, "Câu hỏi thực tiễn chưa có lời đáp định lượng: bền vững hơn BAO NHIÊU — đổi bằng GÌ (thời gian, tiền, độ ổn định)?", size=18, bold=True, color=BLUE)

# --------------------------------------------------------------- Slide 4: gaps
s = add_slide()
title_bar(s, "Khoảng trống nghiên cứu (khảo sát 15 công trình 2023–2025)")
bullets(s, [
    ("①  Chưa có so sánh đối chứng locator-based vs VLM-based trên CÙNG bộ test case, với biến thể giao diện được kiểm soát chủ động — các bài chỉ đo coverage / số bug.", {"size": 18, "space": 16}),
    ("②  Chưa ai đo trực tiếp chi phí bảo trì bộ test (số test / số dòng phải sửa) theo cặp trên cùng một ứng dụng — mới chỉ có chi phí vận hành agent.", {"size": 18, "space": 16}),
    ("③  Khía cạnh bảo mật bỏ ngỏ: (a) PII trên screenshot gửi lên VLM API bên thứ ba; (b) dùng chính VLM kiểm thử phân quyền hiển thị (role-based UI).", {"size": 18, "space": 16}),
], top=Inches(1.6))
tb = box(s, Inches(0.55), Inches(5.9), Inches(12.2), Inches(0.9))
para(tb.text_frame, "Robustness trước thay đổi giao diện chính là LÝ DO VLM được đề xuất — nhưng chưa từng được đo đối chứng.", size=16, color=GREY, first=True)

# ---------------------------------------------------------------- Slide 5: RQs
s = add_slide()
title_bar(s, "Bốn câu hỏi nghiên cứu")
rows = [
    ("RQ1 — Độ bền vững", "Tỉ lệ pass thay đổi thế nào khi giao diện đổi theme / bố cục / biểu tượng?", "← gap ①"),
    ("RQ2 — Chi phí bảo trì", "Số test phải sửa, diff LOC, thời gian phục hồi của hai phương pháp?", "← gap ②"),
    ("RQ3 — Phân quyền hiển thị", "VLM phát hiện lỗi role-based UI chính xác ra sao so với locator?", "← gap ③"),
    ("RQ4 — Che dữ liệu nhạy cảm", "Masking PII trên screenshot ảnh hưởng gì đến độ chính xác định vị?", "← gap ③"),
]
y = 1.55
for name, desc, tag in rows:
    tb = box(s, Inches(0.55), Inches(y), Inches(3.4), Inches(1.1))
    para(tb.text_frame, name, size=18, bold=True, color=BLUE, first=True)
    tb = box(s, Inches(4.05), Inches(y), Inches(7.2), Inches(1.1))
    para(tb.text_frame, desc, size=16, first=True)
    tb = box(s, Inches(11.45), Inches(y), Inches(1.6), Inches(1.1))
    para(tb.text_frame, tag, size=14, color=GREY, first=True)
    y += 1.35

# ------------------------------------------------------------- Slide 6: design
s = add_slide()
title_bar(s, "Thiết kế thực nghiệm: so sánh đối chứng, đối xứng tuyệt đối")
steps = ["App “Mini Shop Manager”\n+ 4 biến thể V0–V3", "2 bộ test song song\n18 test × 2, ánh xạ 1–1",
         "Harness ma trận\n2 pp × 4 biến thể × 5 lặp", "CSV tự động\npass/time/token/$", "Phân tích\n1 lệnh sinh mọi hình/bảng"]
x = 0.45
for i, st in enumerate(steps):
    shp = s.shapes.add_shape(1, Inches(x), Inches(1.7), Inches(2.28), Inches(1.25))
    shp.fill.solid(); shp.fill.fore_color.rgb = BLUE if i % 2 == 0 else ORANGE
    shp.line.fill.background()
    tf = shp.text_frame; tf.word_wrap = True
    for j, ln in enumerate(st.split("\n")):
        para(tf, ln, size=13, bold=(j == 0), color=RGBColor(0xFF, 0xFF, 0xFF),
             align=PP_ALIGN.CENTER, first=(j == 0), space=2)
    x += 2.62
bullets(s, [
    ("Đối xứng: cùng app, cùng 18 kịch bản, cùng viewport 1280×1100, cùng điều kiện chạy — khác duy nhất cách diễn đạt test.", {"size": 17}),
    ("Model ghim: qwen/qwen3-vl-235b-a22b-instruct (OpenRouter) · temperature = 0 · cache TẮT · phiên bản khóa lockfile.", {"size": 17}),
    ("Hai khối mở rộng: RBAC 8 kịch bản × build sạch/cấy lỗi (RQ3) · grounding benchmark 360 call trên ảnh mask (RQ4).", {"size": 17}),
    ("Mọi số liệu ghi tự động vào CSV; toàn bộ repo + số liệu công khai GitHub.", {"size": 17}),
], top=Inches(3.35), height=Inches(3.9))

# ----------------------------------------------------------- Slide 7: variants
s = add_slide()
title_bar(s, "4 biến thể giao diện có kiểm soát — chỉ đổi trình bày, không đổi logic")
labels = [("products-v0.png", "V0 — gốc"), ("products-v1.png", "V1 — dark theme (đổi màu)"),
          ("products-v2.png", "V2 — topbar + đảo cột/icon"), ("products-v3.png", "V3 — đổi icon + nhãn nút")]
pos = [(0.55, 1.3), (6.95, 1.3), (0.55, 4.35), (6.95, 4.35)]
for (fn, lab), (x, y) in zip(labels, pos):
    picture(s, os.path.join(SHOT, fn), Inches(x), Inches(y), width=Inches(3.2))
    tb = box(s, Inches(x + 3.3), Inches(y + 0.9), Inches(2.9), Inches(1.4))
    para(tb.text_frame, lab, size=14, bold=True, color=BLUE if "V0" in lab or "V1" in lab else RED, first=True)
tb = box(s, Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.35))
para(tb.text_frame, "Cấu hình tập trung 1 file (variants.ts) — phủ 3 nguyên nhân gãy test kinh điển: đổi màu / đổi cấu trúc DOM / đổi nhãn.", size=12, color=GREY, first=True)

# -------------------------------------------------------------- Slide 8: RQ1 ★
s = add_slide()
title_bar(s, "RQ1 — Độ bền vững khi giao diện thay đổi", tag="18 test × 5 lặp / biến thể")
picture(s, os.path.join(FIG, "fig-rq1-passrate.png"), Inches(2.97), Inches(1.35), width=Inches(7.4))
big_number(s, "VLM: 18/18 trên CẢ 4 biến thể — Locator: rơi còn 9/18 (V2), 12/18 (V3)",
           "Flakiness = 0 ở cả hai phương pháp (kết quả giống hệt qua 5 lần lặp, temperature = 0)", top=Inches(6.0))

# ----------------------------------------------------------- Slide 9: time/cost
s = add_slide()
title_bar(s, "Giá phải trả: thời gian và chi phí vận hành")
picture(s, os.path.join(FIG, "fig-rq1-time-cost.png"), Inches(1.7), Inches(1.4), width=Inches(9.9))
big_number(s, "VLM ~7–10 phút & $0.05/run   —   Locator ~3 giây & $0/run",
           "Khác biệt thật nằm ở THỜI GIAN, không phải tiền: toàn bộ thực nghiệm hết ≈ $1.5 / ngân sách $50", top=Inches(6.1), color=ORANGE)

# --------------------------------------------------------------- Slide 10: RQ2
s = add_slide()
title_bar(s, "RQ2 — Chi phí bảo trì khi nâng cấp giao diện V0 → Vx")
picture(s, os.path.join(FIG, "fig-rq2-maintenance.png"), Inches(0.85), Inches(1.5), width=Inches(11.6))
big_number(s, "Locator: 15 test / 46 LOC / 314 giây phục hồi  —  VLM: 0 / 0 / 0",
           "Quy trình sửa chuẩn hóa tự động, đo 2 lần độc lập: nội dung sửa trùng khớp từng dòng (tái lập được)", top=Inches(6.1))

# --------------------------------------------------------------- Slide 11: RQ3
s = add_slide()
title_bar(s, "RQ3 — Phát hiện lỗi phân quyền hiển thị (role-based UI)")
picture(s, os.path.join(FIG, "fig-rq3-detection.png"), Inches(1.5), Inches(1.5), width=Inches(10.3))
big_number(s, "Cả hai phương pháp: phát hiện 5/5 lỗi cấy — 0 báo động giả (×5 lặp)",
           "VLM RBAC chỉ ~$0.005/run; kịch bản ngôn ngữ tự nhiên kế thừa luôn độ bền vững của RQ1", top=Inches(6.1), color=GREEN)

# --------------------------------------------------------------- Slide 12: RQ4
s = add_slide()
title_bar(s, "RQ4 — Che PII trước khi gửi VLM API: có mất độ chính xác?")
picture(s, os.path.join(FIG, "fig-rq4-hitrate-iou.png"), Inches(0.55), Inches(1.45), width=Inches(8.6))
picture(s, os.path.join(MASK, "orig", "customer-c01.png"), Inches(9.7), Inches(1.35), width=Inches(2.55))
picture(s, os.path.join(MASK, "masked-pixel", "customer-c01.png"), Inches(9.7), Inches(3.6), width=Inches(2.55))
tb = box(s, Inches(9.4), Inches(5.82), Inches(3.4), Inches(0.32))
para(tb.text_frame, "Trang PII: gốc (trên) / pixelate (dưới)", size=11, color=GREY, first=True)
big_number(s, "Hit-rate 100% ở CẢ 3 điều kiện — IoU chỉ giảm nhẹ tại đúng vùng che",
           "360 lượt gọi (40 phần tử × 3 điều kiện × 3 lặp), kết quả tất định; màn không mask giữ nguyên (đối chứng)", top=Inches(6.15))

# --------------------------------------------------------- Slide 13: trade-off
s = add_slide()
title_bar(s, "Trade-off trung tâm và hàm ý thực tiễn")
hdr = [("Locator (Playwright)", BLUE), ("VLM (Midscene + Qwen3-VL)", ORANGE)]
for i, (h, c) in enumerate(hdr):
    shp = s.shapes.add_shape(1, Inches(0.8 + i * 6.1), Inches(1.5), Inches(5.7), Inches(0.55))
    shp.fill.solid(); shp.fill.fore_color.rgb = c; shp.line.fill.background()
    para(shp.text_frame, h, size=17, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER, first=True)
left_items = ["Chạy ~3 giây, $0", "Gãy 33–50% suite khi đổi bố cục/nhãn", "Phục hồi cần sửa tay: 15 test / 46 LOC", "Trả chi phí MỘT LẦN — khi giao diện đổi"]
right_items = ["Chạy ~8 phút, $0.05/run", "100% pass trên mọi biến thể trình bày", "0 chi phí bảo trì, 0 flakiness (temp 0)", "Trả chi phí MỖI LẦN CHẠY — thời gian + API"]
for i, items in enumerate([left_items, right_items]):
    tb = box(s, Inches(0.8 + i * 6.1), Inches(2.2), Inches(5.7), Inches(2.8))
    first = True
    for it in items:
        para(tb.text_frame, "•  " + it, size=16, first=first, space=10)
        first = False
tb = box(s, Inches(0.8), Inches(5.35), Inches(11.8), Inches(1.6))
para(tb.text_frame, "→ Bổ trợ, không thay thế nhau:", size=19, bold=True, color=BLUE, first=True)
para(tb.text_frame, "CI hai tầng — locator chạy mỗi commit trên giao diện ổn định; VLM làm lớp regression “theo ý định” chạy theo lịch / trước release / giai đoạn giao diện biến động.", size=17)

# ------------------------------------------------------- Slide 14: contributions
s = add_slide()
title_bar(s, "Năm đóng góp (đúng cam kết đề cương)")
contribs = [
    ("C1", "Bộ dữ liệu thực nghiệm tái sử dụng: app + 4 biến thể kiểm soát tập trung (benchmark robustness)"),
    ("C2", "Pipeline đánh giá 3 trục: harness ma trận + giao thức bảo trì chuẩn hóa + script sinh toàn bộ hình/bảng"),
    ("C3", "Bộ số liệu đối chứng đầu tiên (trong phạm vi khảo sát) đo trade-off robustness vs cost trên cùng test case"),
    ("C4", "Cơ chế masking PII + đánh giá định lượng trade-off bảo mật ↔ độ chính xác định vị"),
    ("C5", "Bộ kịch bản kiểm thử phân quyền hiển thị bằng VLM, đối chứng với locator"),
]
y = 1.5
for tag, desc in contribs:
    shp = s.shapes.add_shape(1, Inches(0.55), Inches(y), Inches(0.75), Inches(0.75))
    shp.fill.solid(); shp.fill.fore_color.rgb = BLUE; shp.line.fill.background()
    para(shp.text_frame, tag, size=18, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER, first=True)
    tb = box(s, Inches(1.55), Inches(y + 0.05), Inches(11.2), Inches(0.9))
    para(tb.text_frame, desc, size=16, first=True)
    y += 1.02
tb = box(s, Inches(0.55), Inches(6.75), Inches(12.2), Inches(0.5))
para(tb.text_frame, "Toàn bộ mã nguồn + số liệu công khai: github.com/phanthanh379/doanai — tái lập mọi hình/bảng bằng 1 lệnh.", size=15, color=GREY, first=True)

# ------------------------------------------------ Slide 15: limits & future work
s = add_slide()
title_bar(s, "Hạn chế và hướng phát triển")
tb = box(s, Inches(0.55), Inches(1.5), Inches(5.9), Inches(0.5))
para(tb.text_frame, "Hạn chế (khai báo tường minh)", size=18, bold=True, color=RED, first=True)
bullets(s, [
    "•  1 app mẫu, 1 model, biến thể một chiều có kiểm soát",
    "•  Nhận thức VLM giới hạn trong viewport (1280×1100)",
    "•  Baseline chủ ý dùng selector giòn (worst-case)",
    "•  Thời gian bảo trì đo trên quy trình agent, không phải công người",
], left=Inches(0.55), top=Inches(2.05), width=Inches(5.9), size=15)
tb = box(s, Inches(6.85), Inches(1.5), Inches(5.9), Inches(0.5))
para(tb.text_frame, "Hướng phát triển", size=18, bold=True, color=GREEN, first=True)
bullets(s, [
    "•  Desktop (WinAppDriver) / mobile; trang dài hơn viewport (cuộn + tổng hợp)",
    "•  Thêm model đối chiếu (UI-TARS, self-host) + baseline best-practice (getByRole)",
    "•  Masking tự động (OCR/VLM phát hiện PII); mở rộng lớp lỗi phân quyền",
    "•  CI hai tầng có kiểm soát chi phí",
], left=Inches(6.85), top=Inches(2.05), width=Inches(5.9), size=15)

# ---------------------------------------------------------- Slide 16: conclusion
s = add_slide()
bar = s.shapes.add_shape(1, 0, Inches(2.3), SW, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
tb = box(s, Inches(1.1), Inches(2.7), Inches(11.1), Inches(2.2))
para(tb.text_frame, "Kiểm thử GUI bằng VLM loại bỏ tính giòn trước thay đổi trình bày —", size=24, bold=True, align=PP_ALIGN.CENTER, first=True)
para(tb.text_frame, "100% pass trên mọi biến thể, 0 chi phí bảo trì —", size=24, bold=True, align=PP_ALIGN.CENTER)
para(tb.text_frame, "trả giá bằng thời gian chạy, không phải tiền hay độ ổn định.", size=24, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
tb = box(s, Inches(1.1), Inches(5.3), Inches(11.1), Inches(0.8))
para(tb.text_frame, "Em xin cảm ơn hội đồng — sẵn sàng demo và trả lời câu hỏi.", size=18, color=GREY, align=PP_ALIGN.CENTER, first=True)

# ------------------------------------------------------------ Backup separator
s = add_slide()
tb = box(s, Inches(1.1), Inches(3.2), Inches(11.1), Inches(1.2))
para(tb.text_frame, "Slide dự phòng (Q&A)", size=34, bold=True, color=GREY, align=PP_ALIGN.CENTER, first=True)

# ---------------------------------------------------------------- Backup slides
def backup(title, items):
    s = add_slide()
    title_bar(s, title, tag="backup")
    bullets(s, items, size=16)
    return s

backup("B1 — VLM perception là viewport-bound", [
    "Triệu chứng: aiNumber đếm 9/12 dòng — giống hệt qua mọi lần lặp (deterministic, không phải flaky).",
    "Nguyên nhân: viewport 1280×800 chỉ hiện 9/12 dòng bảng; VLM đếm những gì NÓ THẤY trong screenshot — locator đếm DOM toàn trang.",
    "Xử lý: viewport 1280×1100 cho CẢ HAI suite (giữ đối xứng); phương án “cuộn rồi đếm” bị loại (cuộn làm mất dòng đầu).",
    "Hệ quả: viewport là biến nhiễu phải cố định + báo cáo; kết quả VLM không khái quát sang trang dài hơn viewport nếu thiếu cơ chế cuộn-tổng hợp.",
])
backup("B2 — Quy ước tọa độ grounding của Qwen (0–1000)", [
    "Attempt 1 của RQ4: 340/360 lượt bị chấm “lỗi” → chẩn đoán: VLM định vị ĐÚNG, harness chấm SAI chuẩn tọa độ.",
    "Qwen3-VL luôn trả tọa độ chuẩn hóa 0–1000 bất kể prompt yêu cầu pixel; bằng chứng: box × (1280/1000, 800/1000) trùng ground-truth gần tuyệt đối.",
    "Sửa: prompt theo quy ước gốc của model, harness quy đổi; lưu raw answer (raw-calls.jsonl) để audit; attempt sai giữ làm tư liệu.",
    "Bài học: quy ước của model THẮNG chỉ dẫn trong prompt; phải tách “lỗi model” khỏi “lỗi harness” trước khi kết luận.",
])
backup("B3 — Vì sao Qwen3-VL, không phải GPT-4o/Claude?", [
    "Đề cương ghi “Claude, GPT-4o hoặc tương đương” — khảo sát pilot (GĐ1) cho thấy:",
    "•  GPT-4o: docs Midscene ghi rõ “perform poorly” ở UI localization — chỉ hợp vai planning.",
    "•  Claude: không nằm trong danh sách model Midscene hỗ trợ.",
    "•  Qwen3-VL: nhóm khuyến nghị mặc định về visual grounding; giá $0.20/$0.88 per 1M (rẻ ~10× GPT-4o input); open-weight → nhất quán phương án self-host khi cần bảo mật.",
    "Ghim snapshot 235B-A22B instruct suốt thực nghiệm; fallback: UI-TARS, Qwen2.5-VL.",
])
backup("B4 — Giao thức đo RQ2 (vì sao agent thay vì người bấm giờ?)", [
    "Chỉ dẫn cố định: “sửa tối thiểu để pass lại; chỉ đổi selector/assertion; không refactor; cấm đọc tài liệu kết quả”; mốc T0 → T1 (có danh sách fail) → T2 (18/18).",
    "Mục đích gốc của “cùng một người sửa” là chống thiên lệch giữa 2 bộ — agent cố định thỏa mãn TỐT HƠN: không mệt, không học lỏm, tái lập được.",
    "Kiểm chứng: 2 lần chạy độc lập (agent mới, không ngữ cảnh) → diff trùng khớp TỪNG DÒNG với nhau (9 test/24 LOC ở V2, 6 test/22 LOC ở V3).",
    "2 chỉ số chính (test sửa, diff LOC) bất biến theo người sửa; chỉ số thời gian khai báo hạn chế trong Threats to Validity.",
])
backup("B5 — Baseline có “cố tình yếu” không?", [
    "Lựa chọn thiết kế có chủ ý, khai báo tường minh: baseline đại diện lớp test giòn phổ biến (worst-case của locator).",
    "Nhóm D (icon SVG không aria-label, canvas, div thuần) KHÔNG có lựa chọn ngoài selector cấu trúc — getByRole không cứu được.",
    "getByRole/data-testid giảm rủi ro ở V3 (nhãn) nhưng: data-testid đòi quyền sửa app; và không đổi bản chất kết luận nhóm D.",
    "Hướng phát triển đã ghi: thêm baseline best-practice làm mốc “trường hợp tốt” để kẹp khoảng kết quả.",
])
backup("B6 — Chi phí chi tiết", [
    "RQ1 (suite chính): ~93 AI call/run, ~232k token, $0.050/run — 20 run chính thức ≈ $1.0.",
    "RQ3 (RBAC): 11 call/run, $0.0052/run × 20 run ≈ $0.1.",
    "RQ4 (grounding): 360 call = $0.112.",
    "Tổng toàn thực nghiệm ≈ $1.5 / ngân sách $50 (3%). Giá ghim: $0.20/M in + $0.88/M out.",
    "Token đọc THẬT từ log Midscene từng run (không ước lượng công thức) — cột run_prompt_tokens/run_cost_usd trong CSV.",
])

prs.save(OUT)
print("wrote", OUT, os.path.getsize(OUT) // 1024, "KB,", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
