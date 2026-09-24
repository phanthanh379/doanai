#!/usr/bin/env python3
"""Assemble the thesis draft DOCX from the markdown chapters in docs/bao-cao/.

Output: docs/bao-cao/bao-cao-doan.docx — cover page, auto-updating TOC field,
chapters 1-5 with figures embedded at their [HÌNH x.y] / *Hình 4.x* anchors,
references and appendices. Layout follows common VN thesis conventions
(Times New Roman 13pt, 1.5 line spacing); adjust to the faculty template when
merging. Run from the repo root: python3 harness/make-report-docx.py
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "bao-cao")
OUT = os.path.join(SRC, "bao-cao-doan.docx")
FIG = os.path.join(ROOT, "results", "figures")
SHOT = os.path.join(ROOT, "results", "screenshots-app-v1")

CHAPTER_FILES = [
    "chuong-1-gioi-thieu.md",
    "chuong-2-khao-sat.md",
    "chuong-3-phuong-phap.md",
    "chuong-4-ket-qua.md",
    "chuong-5-ket-luan.md",
    "tai-lieu-tham-khao.md",
    "phu-luc.md",
]

FIG_CAPTIONS = {
    "fig-design-pipeline.png": "Hình 3.1 — Sơ đồ tổng thể thiết kế thực nghiệm",
    "fig-rq1-passrate.png": "Hình 4.1 — RQ1: tỉ lệ pass theo biến thể giao diện (18 test × 5 lặp)",
    "fig-rq1-time-cost.png": "Hình 4.2 — RQ1: thời gian thực thi và chi phí API mỗi run",
    "fig-rq2-maintenance.png": "Hình 4.3 — RQ2: chi phí bảo trì khi nâng cấp giao diện V0 → Vx",
    "fig-rq3-detection.png": "Hình 4.4 — RQ3: phát hiện 5 lỗi phân quyền cấy sẵn, đối chứng build sạch",
    "fig-rq4-hitrate-iou.png": "Hình 4.5 — RQ4: hit-rate và IoU theo điều kiện masking",
    "fig-rq1-rerun-verification.png": "Hình 4.6 — Tái kiểm chứng độc lập RQ1 (19/09/2026): tỉ lệ pass gốc vs. chạy lại",
    "fig-rq2-rerun-verification.png": "Hình 4.7 — Tái kiểm chứng độc lập RQ2 (19/09/2026): diff-LOC theo phong cách sửa và theo phương pháp",
    "fig-rq3-rerun-verification.png": "Hình 4.8 — Tái kiểm chứng độc lập RQ3 (19/09/2026): build sạch vs. build cấy 5 lỗi",
    "fig-rq4-rerun-verification.png": "Hình 4.9 — Tái kiểm chứng độc lập RQ4 (19/09/2026): hit-rate gốc vs. chạy lại (N gấp đôi)",
}

# Danh sách bảng có chú thích (caption) — song song với FIG_CAPTIONS, theo đúng
# thứ tự xuất hiện trong tài liệu. Mỗi mục là dòng chữ đậm được chèn ngay TRÊN
# bảng tương ứng trong markdown nguồn (quy ước VN: caption bảng đặt phía trên).
BANG_CAPTIONS = [
    "Bảng 2.1 — Đối chiếu 27 công trình liên quan trên sáu tiêu chí",
    "Bảng 2.2 — Khoảng trống chiến lược ([19], Table 11) đối chiếu với 4 RQ của đồ án",
    "Bảng 3.1 — Tổng hợp các biến thể giao diện và kịch bản kiểm thử",
    "Bảng 4.1 — RQ1: tỉ lệ pass theo biến thể giao diện",
    "Bảng 4.2 — Phân loại nguyên nhân lỗi (Failure Analysis) áp dụng cho dữ liệu RQ1",
    "Bảng 4.3 — RQ2: chi phí bảo trì theo biến thể (locator vs. VLM)",
    "Bảng 4.4 — RQ3: detection, false alarm và chi phí theo phương pháp",
    "Bảng 4.5 — RQ4: hit-rate và IoU theo màn hình và điều kiện masking",
    "Bảng 4.6 — RQ4: tái kiểm chứng độc lập (N gấp đôi) theo màn hình và điều kiện",
    "Bảng 5.1 — Đối chiếu đóng góp cam kết với sản phẩm thực tế",
]

GREY = RGBColor(0x60, 0x60, 0x60)
TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
CAPTION_NUM_RE = re.compile(r"^(Hình|Bảng)\s+(\d+)\.(\d+)")
_bookmark_counter = [0]


def _bookmark_name(caption_text):
    m = CAPTION_NUM_RE.match(caption_text)
    if not m:
        return None
    kind = "hinh" if m.group(1) == "Hình" else "bang"
    return f"{kind}_{m.group(2)}_{m.group(3)}"


def add_bookmark(paragraph, name):
    _bookmark_counter[0] += 1
    bid = str(_bookmark_counter[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_pageref_entry(doc, label_text, bookmark_name):
    """One line in a Danh mục hình/bảng page: caption text, a tab, and a
    PAGEREF field pointing at the bookmark placed on that caption in the
    body — shows the real page number once the user does Update Field."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(4)
    p.add_run(label_text)
    p.add_run().add_tab()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f"PAGEREF {bookmark_name} \\h")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    for name, size, before in (("Heading 1", 14, 18), ("Heading 2", 13, 14), ("Heading 3", 13, 10)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.top_margin = Cm(3)
        sec.bottom_margin = Cm(3.5)
        sec.left_margin = Cm(3.5)
        sec.right_margin = Cm(2)


def add_page_numbers(section, restart_at_1=False):
    """Add a centered PAGE field to this section's footer. When restart_at_1
    is True, also set the section's page-number-type to restart the Arabic
    sequence at 1 here — used at the section break right before TÓM TẮT, per
    quy định: không đánh số các trang bìa/hội đồng/cảm ơn/mục lục/danh mục."""
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    if restart_at_1:
        sectPr = section._sectPr
        pgNumType = OxmlElement("w:pgNumType")
        pgNumType.set(qn("w:start"), "1")
        sectPr.append(pgNumType)


def add_runs(par, text, base_bold=False):
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        run = par.add_run()
        if tok.startswith("**") and tok.endswith("**"):
            run.text = tok[2:-2]
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run.text = tok[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(11.5)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run.text = tok[1:-1]
            run.italic = True
        else:
            run.text = tok
        if base_bold:
            run.bold = True


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11.5)
    bm = _bookmark_name(text)
    if bm:
        add_bookmark(p, bm)


def add_image(doc, path, width, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=width)
    if caption:
        add_caption(doc, caption)


def add_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GREY


def add_code_block(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(10.5)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for i, r in enumerate(cells) if i != 1]  # drop the |---| separator
    ncols = max(len(r) for r in cells)
    font_size = Pt(10) if ncols >= 7 else Pt(11.5)
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(cells):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(2)
            txt = row[j] if j < len(row) else ""
            add_runs(par, txt, base_bold=(i == 0))
            for run in par.runs:
                run.font.size = font_size
    doc.add_paragraph()


def hinh3_images(doc, text):
    """Replace chapter-3 [HÌNH 3.x ...] placeholders with real screenshots."""
    if text.startswith("[HÌNH 3.2"):
        add_image(doc, os.path.join(SHOT, "products-v0.png"), Inches(5.8),
                  "Hình 3.2 — Giao diện gốc V0, trang Products (vai trò admin)")
        return True
    if text.startswith("[HÌNH 3.3"):
        labels = [("products-v0.png", "V0 — gốc"), ("products-v1.png", "V1 — dark theme"),
                  ("products-v2.png", "V2 — topbar + đảo cột/icon"), ("products-v3.png", "V3 — đổi icon/nhãn")]
        for fn, lab in labels:
            add_image(doc, os.path.join(SHOT, fn), Inches(4.6))
            add_caption(doc, lab)
        add_caption(doc, "Hình 3.3 — Bốn biến thể giao diện có kiểm soát")
        return True
    return False


def render_paragraph(doc, text):
    # Working notes for the drafting process carry no meaning in the deliverable.
    if text.startswith("Ghi chú khi chuyển vào báo cáo"):
        return
    if text.startswith("[HÌNH") or text.startswith("[BẢNG"):
        if hinh3_images(doc, text):
            return
        add_note(doc, "⟪Chỗ chèn: " + text.strip("[]") + "⟫")
        return
    # Table caption lines (bold "Bảng X.Y — ...") get the same bookmarked
    # caption treatment as figures, so Danh mục bảng's PAGEREF resolves.
    bare = text.strip("*").strip()
    if bare in BANG_CAPTIONS:
        add_caption(doc, bare)
        return
    # Figure anchor lines: embed the image + caption, drop the anchor prose
    # (file names and run labels belong to the repo, not the thesis body).
    pngs = [png for png in FIG_CAPTIONS if png in text]
    if pngs:
        for png in pngs:
            add_image(doc, os.path.join(FIG, png), Inches(5.9), FIG_CAPTIONS[png])
        return
    p = doc.add_paragraph()
    add_runs(p, text)


def render_markdown(doc, path):
    lines = open(path, encoding="utf-8").read().splitlines()
    buf, bullets, notes, table, code = [], [], [], [], None
    body_seen = False  # becomes True after the first non-heading body content
    i = 0

    def flush():
        nonlocal buf, bullets, notes, table, body_seen
        if buf:
            render_paragraph(doc, " ".join(buf)); buf = []; body_seen = True
        if bullets:
            for kind, b in bullets:
                if kind == "num":
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.63)
                else:
                    p = doc.add_paragraph(style="List Bullet")
                add_runs(p, b)
            bullets = []; body_seen = True
        if notes:
            text = " ".join(notes); notes = []
            # A quote block before any body content is a drafting note — drop it.
            # Trailing editorial reminders are equally meaningless in the deliverable.
            if body_seen and not text.startswith("Ghi chú khi hoàn thiện"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                add_runs(p, text)
                for run in p.runs:
                    run.italic = True
                body_seen = True
        if table:
            add_table(doc, table); table = []; body_seen = True

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if code is not None:
            if s.startswith("```"):
                add_code_block(doc, code); code = None; body_seen = True
            else:
                code.append(ln)
            i += 1; continue
        if s.startswith("```"):
            flush(); code = []
        elif not s or s == "---":
            flush()
        elif s.startswith("#"):
            flush()
            level = min(len(s) - len(s.lstrip("#")), 3)
            title = s.lstrip("#").strip()
            title = re.sub(r"\s*\((BẢN NHÁP[^)]*|KHUNG CHI TIẾT[^)]*)\)", "", title)
            h = doc.add_heading("", level=level)
            add_runs(h, title)
        elif s.startswith(">"):
            if buf or bullets or table:
                flush()
            notes.append(s.lstrip("> ").strip())
        elif s.startswith("|"):
            if buf or bullets or notes:
                flush()
            table.append(s)
        elif re.match(r"^[-*]\s+", s):
            if buf or table or notes:
                flush()
            bullets.append(("bullet", re.sub(r"^[-*]\s+", "", s)))
        elif re.match(r"^\d+\.\s+", s) and not buf:
            if table or notes:
                flush()
            bullets.append(("num", s))
        elif bullets and ln.startswith(("  ", "\t")):
            bullets[-1] = (bullets[-1][0], bullets[-1][1] + " " + s)
        else:
            if bullets or table or notes:
                flush()
            buf.append(s)
        i += 1
    flush()


def cover_page(doc):
    def center(text, size, bold=False, before=0, after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)

    center("ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH", 14, bold=True)
    center("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN", 14, bold=True)
    center("KHOA [ĐIỀN THEO TEMPLATE]", 13, after=60)
    center("[HỌ TÊN SINH VIÊN 1] — [HỌ TÊN SINH VIÊN 2]", 14, bold=True, before=30)
    center("ĐỒ ÁN TỐT NGHIỆP", 18, bold=True, before=40)
    center("ỨNG DỤNG VISION-LANGUAGE MODEL", 17, bold=True, before=30)
    center("TRONG KIỂM THỬ TỰ ĐỘNG GIAO DIỆN NGƯỜI DÙNG PHẦN MỀM", 17, bold=True)
    center("(Vision-Language Models for Automated Software UI Testing)", 13, after=60)
    center("CỬ NHÂN NGÀNH [TÊN NGÀNH]", 14, bold=True, before=30)
    center("TP. HỒ CHÍ MINH, [NĂM]", 13, before=60)
    doc.add_page_break()


def bia_phu(doc):
    def center(text, size, bold=False, before=0, after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)

    center("ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH", 14, bold=True)
    center("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN", 14, bold=True)
    center("KHOA [ĐIỀN THEO TEMPLATE]", 13, after=30)
    center("[HỌ TÊN SINH VIÊN 1] — [MSSV 1]", 13, before=16)
    center("[HỌ TÊN SINH VIÊN 2] — [MSSV 2]", 13, after=24)
    center("ĐỒ ÁN TỐT NGHIỆP", 18, bold=True, before=20)
    center("ỨNG DỤNG VISION-LANGUAGE MODEL", 17, bold=True, before=20)
    center("TRONG KIỂM THỬ TỰ ĐỘNG GIAO DIỆN NGƯỜI DÙNG PHẦN MỀM", 17, bold=True)
    center("(Vision-Language Models for Automated Software UI Testing)", 13, after=30)
    center("CỬ NHÂN NGÀNH [TÊN NGÀNH]", 14, bold=True, before=16, after=24)
    center("GIẢNG VIÊN HƯỚNG DẪN", 13, bold=True, before=16, after=2)
    center("[HỌ TÊN GIẢNG VIÊN HƯỚNG DẪN]", 13, bold=True, after=30)
    center("TP. HỒ CHÍ MINH, [NĂM]", 13, before=30)
    doc.add_page_break()


def hoi_dong_page(doc):
    h = doc.add_heading("THÔNG TIN HỘI ĐỒNG CHẤM ĐỒ ÁN TỐT NGHIỆP", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Hội đồng chấm Đồ án tốt nghiệp, thành lập theo Quyết định số "
        "[ĐIỀN SỐ QUYẾT ĐỊNH] ngày [ĐIỀN NGÀY] của Hiệu trưởng Trường Đại học "
        "Công nghệ Thông tin."
    )
    for i, role in enumerate(
        ["Chủ tịch", "Thư ký", "Ủy viên phản biện 1", "Ủy viên phản biện 2", "Ủy viên"], start=1
    ):
        doc.add_paragraph(f"{i}. {role} — [ĐIỀN HỌ TÊN]")
    doc.add_page_break()


def loi_cam_on_page(doc):
    doc.add_heading("LỜI CẢM ƠN", level=1)
    doc.add_paragraph(
        "[Điền lời cảm ơn cá nhân — thông thường gồm: Giảng viên hướng dẫn, "
        "quý Thầy/Cô Khoa và Trường đã giảng dạy/hỗ trợ trong suốt quá trình "
        "học tập và thực hiện đồ án, gia đình/bạn bè đã đồng hành. Đoạn này "
        "mang tính cá nhân, không nên để AI viết thay hoàn toàn — nhóm tác giả "
        "tự viết trước khi nộp.]"
    )
    doc.add_page_break()


def danh_muc_hinh_page(doc):
    h = doc.add_heading("DANH MỤC HÌNH VẼ", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for caption in FIG_CAPTIONS.values():
        bm = _bookmark_name(caption)
        add_pageref_entry(doc, caption, bm)
    doc.add_page_break()


def danh_muc_bang_page(doc):
    h = doc.add_heading("DANH MỤC BẢNG", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for caption in BANG_CAPTIONS:
        bm = _bookmark_name(caption)
        add_pageref_entry(doc, caption, bm)
    doc.add_page_break()


ABBREVIATIONS = [
    ("API", "Application Programming Interface"),
    ("CI/CD", "Continuous Integration / Continuous Deployment"),
    ("DOM", "Document Object Model"),
    ("GUI", "Graphical User Interface"),
    ("IoU", "Intersection over Union"),
    ("LLM", "Large Language Model"),
    ("LOC", "Lines of Code"),
    ("PII", "Personally Identifiable Information"),
    ("RBAC", "Role-Based Access Control"),
    ("RQ", "Research Question — Câu hỏi nghiên cứu"),
    ("UI", "User Interface"),
    ("VLM", "Vision-Language Model"),
]


def danh_muc_viet_tat_page(doc):
    h = doc.add_heading("DANH MỤC TỪ VIẾT TẮT", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Từ viết tắt"
    table.rows[0].cells[1].text = "Nguyên nghĩa"
    for abbr, full in ABBREVIATIONS:
        row = table.add_row()
        row.cells[0].text = abbr
        row.cells[1].text = full
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.line_spacing = 1.0
                for run in par.runs:
                    run.font.size = Pt(12)
    doc.add_page_break()


def toc_page(doc):
    h = doc.add_heading("MỤC LỤC", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "(Nhấp chuột phải vào đây rồi chọn Update Field để tạo mục lục.)"
    run.append(t)
    fld.append(run)
    p._p.append(fld)
    doc.add_page_break()


def abstract_page(doc):
    doc.add_heading("TÓM TẮT ĐỒ ÁN", level=1)
    for para in [
        "Đồ án thực hiện một nghiên cứu thực nghiệm so sánh định lượng giữa kiểm thử "
        "GUI dựa trên bộ định vị truyền thống (Playwright) và kiểm thử dựa trên mô hình "
        "thị giác–ngôn ngữ (Midscene.js + Qwen3-VL) trên cùng một ứng dụng web mẫu, cùng "
        "một bộ 18 kịch bản kiểm thử, với bốn biến thể giao diện được kiểm soát chủ động "
        "và mỗi phép đo lặp năm lần.",
        "Kết quả chính: bộ test VLM đạt tỉ lệ pass 100% trên mọi biến thể giao diện và "
        "không cần sửa dòng mã nào (chi phí bảo trì bằng 0), trong khi bộ test locator rơi "
        "còn 50–67% trên các biến thể đổi cấu trúc/nhãn và cần sửa 15 lượt test / 46 dòng "
        "mã để phục hồi; đổi lại, mỗi lượt chạy VLM tốn ~7–10 phút và ~0,05 USD so với ~3 "
        "giây và 0 USD của locator, với độ ổn định tuyệt đối ở cả hai phương pháp. Hai "
        "khảo sát mở rộng cho thấy VLM phát hiện đủ 5/5 lỗi phân quyền hiển thị cấy sẵn "
        "không báo động giả, và việc che dữ liệu nhạy cảm trên ảnh chụp màn hình không làm "
        "giảm khả năng định vị phần tử của mô hình.",
        "Toàn bộ ứng dụng mẫu, hai bộ kiểm thử, hạ tầng đo lường và số liệu được công khai "
        "để tái lập.",
    ]:
        doc.add_paragraph(para)
    doc.add_page_break()


def main():
    doc = Document()
    base_styles(doc)
    # Phần đầu (bìa, hội đồng, cảm ơn, mục lục, danh mục) — KHÔNG đánh số trang.
    cover_page(doc)
    bia_phu(doc)
    hoi_dong_page(doc)
    loi_cam_on_page(doc)
    toc_page(doc)
    danh_muc_hinh_page(doc)
    danh_muc_bang_page(doc)
    danh_muc_viet_tat_page(doc)
    # Cắt section mới ngay trước TÓM TẮT: đây là nơi đánh số trang bắt đầu (số 1).
    numbered_section = doc.add_section(WD_SECTION.NEW_PAGE)
    numbered_section.top_margin = Cm(3)
    numbered_section.bottom_margin = Cm(3.5)
    numbered_section.left_margin = Cm(3.5)
    numbered_section.right_margin = Cm(2)
    add_page_numbers(numbered_section, restart_at_1=True)
    abstract_page(doc)
    for i, fn in enumerate(CHAPTER_FILES):
        render_markdown(doc, os.path.join(SRC, fn))
        if i < len(CHAPTER_FILES) - 1:
            doc.add_page_break()
    doc.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT) // 1024, "KB")


if __name__ == "__main__":
    main()
